"""
文档解析器。

支持格式：.md / .txt / .docx / .pdf / .pptx

所有解析器返回统一的 DocumentParseResult。
"""

import os
import re
from typing import Optional

from .schema import (
    DocumentParseResult,
    DocumentSection,
    DocumentParseError,
    DocumentUnsupportedError,
)


# 支持的扩展名 → format 名映射
_EXT_MAP = {
    ".md": "md",
    ".txt": "txt",
    ".docx": "docx",
    ".pdf": "pdf",
    ".pptx": "pptx",
}


def parse_document(path: str, doc_id: Optional[str] = None) -> DocumentParseResult:
    """解析文档，返回统一结构。

    Args:
        path: 文档文件路径。
        doc_id: 文档 ID，默认用文件名（不含扩展名）。

    Returns:
        DocumentParseResult，其 status 字段指示解析状态。
    """
    if not os.path.exists(path):
        raise DocumentParseError(f"文件不存在: {path}")

    ext = os.path.splitext(path)[1].lower()
    if ext not in _EXT_MAP:
        raise DocumentUnsupportedError(f"不支持的文件格式: {ext}")

    filename = os.path.basename(path)
    if doc_id is None:
        doc_id = os.path.splitext(filename)[0]

    fmt = _EXT_MAP[ext]

    try:
        if fmt == "md":
            return _parse_md(path, doc_id, filename)
        elif fmt == "txt":
            return _parse_txt(path, doc_id, filename)
        elif fmt == "docx":
            return _parse_docx(path, doc_id, filename)
        elif fmt == "pdf":
            return _parse_pdf(path, doc_id, filename)
        elif fmt == "pptx":
            return _parse_pptx(path, doc_id, filename)
    except Exception as exc:
        return DocumentParseResult(
            doc_id=doc_id,
            filename=filename,
            format=fmt,
            title=filename,
            pages=0,
            text="",
            status="failed",
            error=str(exc),
        )


# ====================================================================
# Markdown
# ====================================================================

def _parse_md(path: str, doc_id: str, filename: str) -> DocumentParseResult:
    with open(path, "r", encoding="utf-8") as f:
        content = f.read()

    lines = content.split("\n")
    sections: list[DocumentSection] = []
    heading_stack: list[str] = []
    current_title = ""
    current_lines: list[str] = []
    in_code_block = False
    section_counter = 0

    def _flush_section():
        nonlocal section_counter
        if not current_lines and not current_title:
            return
        text = "\n".join(current_lines).strip()
        if not text and not current_title:
            return
        section_counter += 1
        sections.append(DocumentSection(
            section_id=f"sec_{section_counter:04d}",
            title=current_title,
            heading_path=list(heading_stack),
            page_start=None,
            page_end=None,
            text=text,
        ))

    for line in lines:
        # 代码块检测
        if line.strip().startswith("```"):
            in_code_block = not in_code_block
            current_lines.append(line)
            continue

        if in_code_block:
            current_lines.append(line)
            continue

        # 标题检测
        heading_match = re.match(r'^(#{1,6})\s+(.+)$', line)
        if heading_match:
            _flush_section()
            level = len(heading_match.group(1))
            title_text = heading_match.group(2).strip()

            # 维护 heading_stack
            # 去掉 >= 当前 level 的标题
            while len(heading_stack) >= level:
                heading_stack.pop()
            heading_stack.append(title_text)

            current_title = title_text
            current_lines = []
            continue

        current_lines.append(line)

    _flush_section()

    title = sections[0].title if sections else filename
    full_text = "\n".join(line for line in lines)

    status, warnings = _check_result(sections, full_text)

    return DocumentParseResult(
        doc_id=doc_id,
        filename=filename,
        format="md",
        title=title,
        pages=1,
        text=full_text,
        sections=sections,
        status=status,
        warnings=warnings,
    )


# ====================================================================
# TXT
# ====================================================================

def _parse_txt(path: str, doc_id: str, filename: str) -> DocumentParseResult:
    text = _read_with_encoding(path)

    # 按双换行切段落组
    raw_groups = re.split(r'\n\s*\n', text)
    groups = [g.strip() for g in raw_groups if g.strip()]

    # 合并过短组（< 100 字）
    merged: list[str] = []
    buf = ""
    for g in groups:
        if len(buf) + len(g) < 100:
            buf = (buf + "\n\n" + g).strip()
        else:
            if buf:
                merged.append(buf)
            buf = g
    if buf:
        merged.append(buf)

    sections: list[DocumentSection] = []
    for i, g in enumerate(merged):
        # 尝试检测第一行是否为标题（短行 + 后续内容长）
        lines = g.split("\n")
        sec_title = ""
        sec_text = g
        if len(lines) > 1 and len(lines[0]) <= 40:
            sec_title = lines[0].strip()
            sec_text = "\n".join(lines[1:]).strip()

        sections.append(DocumentSection(
            section_id=f"sec_{i + 1:04d}",
            title=sec_title,
            heading_path=[sec_title] if sec_title else [],
            page_start=None,
            page_end=None,
            text=sec_text,
        ))

    title = sections[0].title if sections else filename
    status, warnings = _check_result(sections, text)

    return DocumentParseResult(
        doc_id=doc_id,
        filename=filename,
        format="txt",
        title=title,
        pages=1,
        text=text,
        sections=sections,
        status=status,
        warnings=warnings,
    )


# ====================================================================
# DOCX
# ====================================================================

def _parse_docx(path: str, doc_id: str, filename: str) -> DocumentParseResult:
    import docx

    doc = docx.Document(path)

    sections: list[DocumentSection] = []
    heading_stack: list[str] = []
    current_title = ""
    current_lines: list[str] = []
    section_counter = 0
    full_lines: list[str] = []

    def _flush():
        nonlocal section_counter
        text = "\n".join(current_lines).strip()
        if not text and not current_title:
            return
        section_counter += 1
        sections.append(DocumentSection(
            section_id=f"sec_{section_counter:04d}",
            title=current_title,
            heading_path=list(heading_stack),
            page_start=None,
            page_end=None,
            text=text,
        ))

    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        para_text = para.text.strip()
        full_lines.append(para.text)

        # 检测 Word 样式标题
        heading_level = _detect_heading_level(style_name, para_text)
        if heading_level is not None:
            _flush()

            while len(heading_stack) >= heading_level:
                heading_stack.pop()
            heading_stack.append(para_text)

            current_title = para_text
            current_lines = []
            continue

        if para_text:
            current_lines.append(para_text)

    _flush()

    # 提取表格文本
    for table in doc.tables:
        table_lines: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_lines.append(" | ".join(cells))
        if table_lines:
            section_counter += 1
            sections.append(DocumentSection(
                section_id=f"sec_{section_counter:04d}",
                title="表格",
                heading_path=[],
                page_start=None,
                page_end=None,
                text="\n".join(table_lines),
            ))

    title = sections[0].title if sections else filename
    full_text = "\n".join(full_lines)
    status, warnings = _check_result(sections, full_text)

    return DocumentParseResult(
        doc_id=doc_id,
        filename=filename,
        format="docx",
        title=title,
        pages=1,
        text=full_text,
        sections=sections,
        status=status,
        warnings=warnings,
    )


def _detect_heading_level(style_name: str, text: str) -> Optional[int]:
    """检测段落是否为标题，返回标题层级（1-based），否则 None。"""
    # 方式 1：Word 样式
    style_lower = style_name.lower()
    for level in range(1, 7):
        if f"heading {level}" in style_lower or f"heading{level}" in style_lower:
            return level

    # 方式 2：启发式 —— 短行 + 编号模式
    if len(text) <= 40:
        if re.match(r'^第[一二三四五六七八九十\d]+[章节]', text):
            return 2
        if re.match(r'^\d+[\.\s]', text):
            return 3
        if re.match(r'^\d+\.\d+', text):
            return 4
        if re.match(r'^[（(][一二三四五六七八九十\d]+[）)]', text):
            return 3
        if re.match(r'^[一二三四五六七八九十]、', text):
            return 3

    return None


# ====================================================================
# PDF
# ====================================================================

def _parse_pdf(path: str, doc_id: str, filename: str) -> DocumentParseResult:
    from pypdf import PdfReader

    reader = PdfReader(path)
    all_pages_text: list[str] = []
    skipped_pages: list[int] = []
    warnings: list[str] = []

    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text()
            if text:
                all_pages_text.append(text.strip())
            else:
                all_pages_text.append("")
                skipped_pages.append(i + 1)
        except Exception:
            all_pages_text.append("")
            skipped_pages.append(i + 1)

    if skipped_pages:
        warnings.append(f"跳过空页: {skipped_pages}")

    total_pages = len(reader.pages)

    # 按页合并短页（< 120 字，最多连续 3 页）
    merged_pages: list[dict] = []  # [{page_start, page_end, text}]
    buf_pages: list[int] = []
    buf_text: list[str] = []
    buf_len = 0

    def _flush_merged():
        nonlocal buf_pages, buf_text, buf_len
        if buf_pages:
            merged_pages.append({
                "page_start": buf_pages[0],
                "page_end": buf_pages[-1],
                "text": "\n".join(buf_text),
            })
            buf_pages = []
            buf_text = []
            buf_len = 0

    for i, text in enumerate(all_pages_text):
        page_num = i + 1
        if len(text) < 120 and len(buf_pages) < 3:
            buf_pages.append(page_num)
            buf_text.append(text)
            buf_len += len(text)
            if buf_len >= 300:
                _flush_merged()
        else:
            _flush_merged()
            buf_pages.append(page_num)
            buf_text.append(text)
            _flush_merged()

    _flush_merged()

    # 将合并后的页面组转为 sections
    sections: list[DocumentSection] = []
    for idx, mp in enumerate(merged_pages):
        if not mp["text"].strip():
            continue

        # 在页内按空行切段落组
        groups = re.split(r'\n\s*\n', mp["text"])
        groups = [g.strip() for g in groups if g.strip()]

        # 合并短组（< 100 字）
        merged_groups: list[str] = []
        buf = ""
        for g in groups:
            clean = re.sub(r'\s+', ' ', g).strip()
            if len(buf) + len(clean) < 100:
                buf = (buf + "\n\n" + clean).strip()
            else:
                if buf:
                    merged_groups.append(buf)
                buf = clean
        if buf:
            merged_groups.append(buf)

        for g_idx, g in enumerate(merged_groups):
            # 给 section 命名
            if mp["page_start"] == mp["page_end"]:
                sec_title = f"第 {mp['page_start']} 页"
            else:
                sec_title = f"第 {mp['page_start']}-{mp['page_end']} 页"

            sections.append(DocumentSection(
                section_id=f"sec_{idx + 1:04d}_{g_idx:02d}",
                title=sec_title,
                heading_path=[sec_title],
                page_start=mp["page_start"],
                page_end=mp["page_end"],
                text=g,
            ))

    full_text = "\n".join(all_pages_text)
    title = filename
    status, check_warnings = _check_result(sections, full_text)
    warnings.extend(check_warnings)

    if not full_text.strip():
        status = "empty_text"

    return DocumentParseResult(
        doc_id=doc_id,
        filename=filename,
        format="pdf",
        title=title,
        pages=total_pages,
        text=full_text,
        sections=sections,
        status=status,
        warnings=warnings,
    )


# ====================================================================
# PPTX
# ====================================================================

def _parse_pptx(path: str, doc_id: str, filename: str) -> DocumentParseResult:
    from pptx import Presentation

    prs = Presentation(path)
    sections: list[DocumentSection] = []
    all_texts: list[str] = []

    for i, slide in enumerate(prs.slides):
        slide_num = i + 1
        shape_texts: list[str] = []
        slide_title = ""

        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        shape_texts.append(t)
                        # 第一个 shape 的第一段文字视为标题
                        if not slide_title:
                            slide_title = t

        slide_text = "\n".join(shape_texts)
        all_texts.append(slide_text)

        if not slide_text.strip():
            continue

        sections.append(DocumentSection(
            section_id=f"sec_{slide_num:04d}",
            title=slide_title or f"幻灯片 {slide_num}",
            heading_path=[slide_title] if slide_title else [],
            page_start=slide_num,
            page_end=slide_num,
            text=slide_text,
        ))

    title = sections[0].title if sections else filename
    full_text = "\n\n".join(all_texts)
    status, warnings = _check_result(sections, full_text)

    return DocumentParseResult(
        doc_id=doc_id,
        filename=filename,
        format="pptx",
        title=title,
        pages=len(prs.slides),
        text=full_text,
        sections=sections,
        status=status,
        warnings=warnings,
    )


# ====================================================================
# 通用辅助
# ====================================================================

def _read_with_encoding(path: str) -> str:
    """按优先级尝试编码读取文本文件。"""
    encodings = ["utf-8-sig", "utf-8", "gbk", "gb2312", "utf-16"]
    for enc in encodings:
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    # 最终兜底
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _check_result(
    sections: list[DocumentSection], full_text: str
) -> tuple[str, list[str]]:
    """根据 sections 和全文判断解析状态。"""
    warnings: list[str] = []

    if not sections:
        return ("empty_text", ["文档无有效内容"])

    if not full_text.strip():
        return ("empty_text", warnings)

    # 检查是否有 short_section
    for sec in sections:
        if len(sec.text) < 50 and sec.text.strip():
            warnings.append(f"{sec.section_id} ({sec.title}) 内容过短 ({len(sec.text)} 字)")

    return ("ok", warnings)
